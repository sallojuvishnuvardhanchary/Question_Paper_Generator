"""
Syllabus Parser Service for CoreAlgorithm PROBLEM95.
Extracts raw text content from uploaded PDF, DOCX, and TXT syllabus documents.
"""

import os
from typing import Dict, Any
import pypdf
import docx

def parse_uploaded_syllabus(file_path: str, filename: str) -> str:
    """
    Parses and extracts readable text from an uploaded document.
    Supported extensions: .pdf, .docx, .txt
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found on disk: {file_path}")

    ext = os.path.splitext(filename.lower())[1]

    if ext == '.pdf':
        return _extract_from_pdf(file_path)
    elif ext == '.docx':
        return _extract_from_docx(file_path)
    elif ext in ('.txt', '.text'):
        return _extract_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Please upload a PDF, DOCX, or TXT file.")

def _extract_from_pdf(file_path: str) -> str:
    """Extracts text page by page from a PDF file."""
    text_chunks = []
    try:
        reader = pypdf.PdfReader(file_path)
        if len(reader.pages) == 0:
            raise ValueError("The uploaded PDF has 0 pages.")

        for page_idx, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text.strip())

        extracted = "\n\n".join(text_chunks).strip()
        if not extracted:
            raise ValueError(
                "Could not extract readable text from the uploaded PDF. "
                "The PDF might be a scanned image or empty. Please upload a text-based document or TXT copy."
            )
        return extracted
    except Exception as e:
        if isinstance(e, ValueError):
            raise e
        raise ValueError(f"Failed to parse PDF document: {str(e)}")

def _extract_from_docx(file_path: str) -> str:
    """Extracts text from paragraphs and tables in a DOCX file."""
    try:
        doc = docx.Document(file_path)
        lines = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                lines.append(text)

        # Extract table cells if any
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    lines.append(" | ".join(row_texts))

        extracted = "\n".join(lines).strip()
        if not extracted:
            raise ValueError("The uploaded DOCX file appears to be empty.")
        return extracted
    except Exception as e:
        if isinstance(e, ValueError):
            raise e
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")

def _extract_from_txt(file_path: str) -> str:
    """Extracts text directly from a plain text file."""
    encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']
    for enc in encodings:
        try:
            with open(file_path, 'r', encoding=enc) as f:
                content = f.read().strip()
                if content:
                    return content
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode TXT file using standard text encodings.")
