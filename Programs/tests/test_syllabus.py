"""
Unit & Integration Tests for Syllabus Parsing & Structure Extraction
CoreAlgorithm PROBLEM95 - Syllabus-Driven Exam Question Paper Generator
"""

import os
import pytest
import tempfile
from services.syllabus_parser import parse_uploaded_syllabus
from services.syllabus_structure import extract_syllabus_structure

def test_parse_txt_syllabus():
    sample_text = """
    Operating Systems Syllabus
    UNIT I: Introduction to OS
    Process Concepts, Process Scheduling, Inter-process Communication.
    
    UNIT II: Synchronization & Deadlocks
    Critical Section Problem, Semaphores, Deadlock Characterization.
    """
    with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False, encoding='utf-8') as tf:
        tf.write(sample_text)
        temp_path = tf.name

    try:
        parsed = parse_uploaded_syllabus(temp_path, "os_syllabus.txt")
        assert "Operating Systems" in parsed
        assert "UNIT I" in parsed
        assert "UNIT II" in parsed
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_parse_docx_syllabus():
    from docx import Document
    doc = Document()
    doc.add_heading('Database Management Systems', level=1)
    doc.add_heading('UNIT I: Relational Model', level=2)
    doc.add_paragraph('Entity Relationship Diagrams, Relational Algebra, Relational Calculus.')
    doc.add_heading('UNIT II: SQL & Normalization', level=2)
    doc.add_paragraph('DDL, DML, First Normal Form, 2NF, 3NF, BCNF.')
    
    with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as tf:
        temp_path = tf.name
    doc.save(temp_path)

    try:
        parsed = parse_uploaded_syllabus(temp_path, "dbms.docx")
        assert "Database Management Systems" in parsed
        assert "Relational Model" in parsed
        assert "Normalization" in parsed
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_parse_pdf_syllabus():
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tf:
        temp_path = tf.name

    c = canvas.Canvas(temp_path, pagesize=letter)
    c.drawString(100, 750, "Computer Networks (CS-502)")
    c.drawString(100, 720, "UNIT I: Physical and Data Link Layer")
    c.drawString(100, 700, "OSI Reference Model, TCP/IP, Framing, Error Detection.")
    c.drawString(100, 670, "UNIT II: Network Layer")
    c.drawString(100, 650, "IPv4, IPv6, Routing Algorithms, Dijkstra, Distance Vector.")
    c.save()

    try:
        parsed = parse_uploaded_syllabus(temp_path, "networks.pdf")
        assert "Computer Networks" in parsed
        assert "Physical" in parsed
        assert "Routing Algorithms" in parsed
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_unsupported_extension_error():
    with tempfile.NamedTemporaryFile(suffix='.exe', delete=False) as tf:
        temp_path = tf.name
    try:
        with pytest.raises(ValueError, match="Unsupported file format"):
            parse_uploaded_syllabus(temp_path, "test.exe")
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_extract_syllabus_structure_roman_units():
    syllabus_raw = """
    DESIGN AND ANALYSIS OF ALGORITHMS
    
    UNIT I: Algorithm Analysis and Divide & Conquer
    Asymptotic Notation, Recurrence Relations, Master Theorem, Binary Search, Merge Sort, Quick Sort.
    
    UNIT II: Disjoint Sets and Greedy Method
    Disjoint Set Operations, Union and Find, Minimum Spanning Trees, Prim's, Kruskal's, Dijkstra's Algorithm, Huffman Coding.
    
    UNIT III: Dynamic Programming
    0/1 Knapsack, All Pairs Shortest Path, Floyd-Warshall, Matrix Chain Multiplication, Longest Common Subsequence.
    
    UNIT IV: Backtracking and Branch & Bound
    8-Queens Problem, Sum of Subsets, Graph Coloring, Hamiltonian Cycles, Travelling Salesperson Problem.
    
    UNIT V: NP-Completeness
    P and NP classes, NP-Hard, NP-Complete, Cook's Theorem, Vertex Cover, Approximation Algorithms.
    """
    struct = extract_syllabus_structure(syllabus_raw)
    assert struct['subject'] == 'DESIGN AND ANALYSIS OF ALGORITHMS'
    assert len(struct['units']) == 5

    # Check Unit 1
    u1 = struct['units'][0]
    assert u1['unit'] == 1
    assert "Divide & Conquer" in u1['title']
    assert any("Merge Sort" in t for t in u1['topics'])

    # Check Unit 3
    u3 = struct['units'][2]
    assert u3['unit'] == 3
    assert any("Knapsack" in t for t in u3['topics'])

    # Check Unit 5
    u5 = struct['units'][4]
    assert u5['unit'] == 5
    assert any("Cook's Theorem" in t or "NP" in t for t in u5['topics'])

def test_extract_unstructured_fallback():
    raw_text = "Overview of Software Engineering. Requirements Gathering. Architecture Design. Testing and Quality Assurance. Agile Scrum."
    struct = extract_syllabus_structure(raw_text, default_subject="Software Engineering")
    assert struct['subject'] == "Software Engineering"
    assert len(struct['units']) >= 1
    total_topics = sum(len(u.get('topics', [])) for u in struct['units'])
    assert total_topics > 0
